import io
import logging
from typing import Optional

import docx          # python-docx
import pymupdf       # PyMuPDF
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.document import Document, DocumentChunk
from app.services.embedding_service import embed_texts
from app.utils.file_validation import compute_hash, get_extension
from app.utils.text_utils import clean_text, split_text

logger = logging.getLogger(__name__)


class DuplicateDocumentError(Exception):
    """Raised when the same file was already uploaded."""

    def __init__(self, document: Document):
        self.document = document
        super().__init__("Duplicate document")


# Each "page" is a tuple: (page_number or None, text)
PageText = tuple[Optional[int], str]


def _extract_pdf(content: bytes) -> list[PageText]:
    try:
        pdf = pymupdf.open(stream=content, filetype="pdf")
    except Exception:
        raise ValueError("Could not open this PDF. The file may be damaged.")

    with pdf:
        if pdf.needs_pass:
            raise ValueError("This PDF is password protected.")
        pages = []
        for number, page in enumerate(pdf, start=1):
            text = clean_text(page.get_text())
            if text:
                pages.append((number, text))
    return pages


def _extract_docx(content: bytes) -> list[PageText]:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception:
        raise ValueError("Could not open this DOCX file. The file may be damaged.")

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Also read text inside tables
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = clean_text("\n".join(parts))
    return [(None, text)] if text else []   # DOCX has no fixed page numbers


def _extract_txt(content: bytes) -> list[PageText]:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("latin-1")   # fallback for older files
    text = clean_text(text)
    return [(None, text)] if text else []


def extract_pages(content: bytes, extension: str) -> list[PageText]:
    """Pick the right reader based on the file type."""
    if extension == ".pdf":
        return _extract_pdf(content)
    if extension == ".docx":
        return _extract_docx(content)
    return _extract_txt(content)


def process_document(
    db: Session, filename: str, content: bytes, category: str = "General"
) -> Document:
    """Full pipeline: save -> extract -> chunk -> embed -> store."""
    extension = get_extension(filename)
    file_hash = compute_hash(content)

    # Step 1: check for duplicates
    existing = db.execute(
        select(Document).where(Document.file_hash == file_hash)
    ).scalar_one_or_none()

    if existing:
        if existing.status == "failed":
            # An earlier try failed, so remove it and try again
            db.delete(existing)
            db.commit()
        else:
            raise DuplicateDocumentError(existing)

    # Step 2: save the document record with status "processing"
    document = Document(
        filename=filename,
        file_type=extension.lstrip("."),
        file_size=len(content),
        file_hash=file_hash,
        category=category,
        status="processing",
        file_data=content,   # original file stored in Neon (MVP storage)
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    document_id = document.id

    try:
        # Step 3: extract text
        pages = extract_pages(content, extension)
        if not pages:
            raise ValueError(
                "No readable text found. Scanned or image-only files are not supported."
            )

        # Step 4: split each page into chunks (keep the page number)
        pieces: list[tuple[Optional[int], str]] = []
        for page_number, page_text in pages:
            for piece in split_text(
                page_text, settings.CHUNK_SIZE, settings.CHUNK_OVERLAP
            ):
                pieces.append((page_number, piece))

        # Step 5: create embeddings for all chunks
        vectors = embed_texts([text for _, text in pieces])

        # Step 6: store chunks
        chunk_rows = [
            DocumentChunk(
                document_id=document_id,
                chunk_index=index,
                page_number=page_number,
                content=text,
                embedding=vector,
            )
            for index, ((page_number, text), vector) in enumerate(zip(pieces, vectors))
        ]
        db.add_all(chunk_rows)

        # Step 7: mark the document as ready
        document.extracted_text = "\n\n".join(text for _, text in pages)
        document.page_count = len(pages) if extension == ".pdf" else None
        document.chunk_count = len(pieces)
        document.status = "ready"
        db.commit()
        logger.info("Document %s processed: %s chunks", document_id, len(pieces))
        return document

    except Exception as error:
        # Undo partial work and mark the document as failed
        db.rollback()
        failed = db.get(Document, document_id)
        if failed:
            failed.status = "failed"
            failed.error_message = str(error)[:500]
            db.commit()
        logger.error("Processing failed for document %s: %s", document_id, error)
        raise